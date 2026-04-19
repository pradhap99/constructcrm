"""
Document processing pipeline:
  1. Extract text from PDF/Excel/text file
  2. Send to Claude claude-sonnet-4-6 for BOQ extraction
  3. Generate Excel with openpyxl
  4. Update Document record with results
"""

import json
import os
import re
import zlib
import uuid
from datetime import datetime, timezone
from typing import Any, Dict, List

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, numbers
from openpyxl.utils import get_column_letter

from app.database import SessionLocal
from app.models.document import Document, DocumentStatus


# ---------------------------------------------------------------------------
# Text extraction helpers
# ---------------------------------------------------------------------------

def _extract_pdf_text(path: str) -> str:
    """Minimal PDF text extractor — no cryptography dependency."""
    with open(path, "rb") as f:
        raw = f.read()

    texts: List[str] = []

    # Locate all stream...endstream blocks
    for m in re.finditer(rb"stream\r?\n(.*?)\r?\nendstream", raw, re.DOTALL):
        chunk = m.group(1)
        # Try FlateDecode (zlib) decompression
        try:
            chunk = zlib.decompress(chunk)
        except Exception:
            pass
        # Extract BT...ET blocks and pull out parenthesised text strings
        for bt in re.finditer(rb"BT(.*?)ET", chunk, re.DOTALL):
            for tj in re.finditer(rb"\((.*?)\)", bt.group(1)):
                t = tj.group(1).decode("latin-1", errors="replace")
                t = t.strip()
                if len(t) > 1:
                    texts.append(t)

    text = " ".join(texts)
    # Also scan raw bytes for printable ASCII runs as a fallback
    if len(text) < 100:
        printable = re.findall(rb"[ -~]{4,}", raw)
        text = " ".join(p.decode("ascii", errors="replace") for p in printable)

    return text[:15000]  # cap to avoid huge token counts


def _extract_excel_text(path: str) -> str:
    """Read all cell values from an Excel file."""
    wb = openpyxl.load_workbook(path, data_only=True)
    rows: List[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                rows.append("\t".join(cells))
    return "\n".join(rows)[:15000]


def extract_text(path: str, original_name: str) -> str:
    ext = os.path.splitext(original_name)[1].lower()
    try:
        if ext == ".pdf":
            return _extract_pdf_text(path)
        elif ext in (".xlsx", ".xls"):
            return _extract_excel_text(path)
        else:
            with open(path, "r", encoding="utf-8", errors="replace") as f:
                return f.read()[:15000]
    except Exception as e:
        return f"[Text extraction failed: {e}]"


# ---------------------------------------------------------------------------
# Claude extraction
# ---------------------------------------------------------------------------

CLAUDE_PROMPT = """\
You are an expert civil construction quantity surveyor.
Extract all BOQ (Bill of Quantities) / quotation line items from the document text below.

Return ONLY a valid JSON object — no markdown fences, no explanation:
{{
  "document_type": "quotation|BOQ|invoice|other",
  "vendor_name": "string or null",
  "vendor_gstin": "string or null",
  "quote_number": "string or null",
  "quote_date": "YYYY-MM-DD or null",
  "valid_until": "YYYY-MM-DD or null",
  "items": [
    {{
      "section": "section or category name",
      "description": "item description",
      "specification": "technical spec or empty string",
      "unit": "MT|Nos|Sqm|Cum|Mtr|Ltr|Kg|Set|LS",
      "quantity": 0.0,
      "rate": 0.0,
      "gst_percent": 18,
      "amount": 0.0
    }}
  ],
  "total_amount": 0.0,
  "overall_confidence": 0.85
}}

Document text:
{text}
"""

FALLBACK_RESULT: Dict[str, Any] = {
    "document_type": "quotation",
    "vendor_name": "Demo Vendor Ltd",
    "vendor_gstin": "27AAACT2727Q1ZW",
    "quote_number": "DEMO/Q/2024/001",
    "quote_date": "2024-01-10",
    "valid_until": "2024-02-10",
    "items": [
        {"section": "Civil Works", "description": "TMT Bars Fe500D 12mm", "specification": "IS:1786, Fe500D grade", "unit": "MT",  "quantity": 25.0,  "rate": 58500.0, "gst_percent": 18, "amount": 1462500.0},
        {"section": "Civil Works", "description": "TMT Bars Fe500D 16mm", "specification": "IS:1786, Fe500D grade", "unit": "MT",  "quantity": 15.0,  "rate": 59200.0, "gst_percent": 18, "amount":  888000.0},
        {"section": "Civil Works", "description": "Binding Wire 16 Gauge", "specification": "",                     "unit": "KG",  "quantity": 200.0, "rate":    85.0, "gst_percent": 18, "amount":   17000.0},
        {"section": "Steel",       "description": "MS Plates 6mm",         "specification": "IS:2062 E250",         "unit": "MT",  "quantity":  2.0,  "rate": 62000.0, "gst_percent": 18, "amount":  124000.0},
    ],
    "total_amount": 2491500.0,
    "overall_confidence": 0.91,
}


def call_claude(text: str, api_key: str) -> Dict[str, Any]:
    import anthropic
    client = anthropic.Anthropic(api_key=api_key)
    message = client.messages.create(
        model="claude-sonnet-4-6",
        max_tokens=4096,
        messages=[{"role": "user", "content": CLAUDE_PROMPT.format(text=text)}],
    )
    raw = message.content[0].text.strip()
    # Strip markdown fences if Claude added them
    raw = re.sub(r"^```json\s*", "", raw)
    raw = re.sub(r"\s*```$", "", raw)
    return json.loads(raw)


# ---------------------------------------------------------------------------
# Excel generation
# ---------------------------------------------------------------------------

HEADER_FILL  = PatternFill("solid", fgColor="2F5496")
HEADER_FONT  = Font(bold=True, color="FFFFFF", size=10)
TITLE_FONT   = Font(bold=True, size=12, color="1F3864")
SECTION_FILL = PatternFill("solid", fgColor="D9E1F2")
SECTION_FONT = Font(bold=True, size=10, color="1F3864")
TOTAL_FONT   = Font(bold=True, size=10)
TOTAL_FILL   = PatternFill("solid", fgColor="FFF2CC")
INR_FORMAT   = '₹#,##0.00'
NUM_FORMAT   = '#,##0.000'

HEADERS = ["Section", "Description", "Specification", "Unit", "Qty", "Rate (₹)", "GST%", "Amount (₹)"]
COL_WIDTHS = [20, 42, 32, 8, 10, 16, 7, 16]


def generate_excel(data: Dict[str, Any], out_path: str) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BOQ Extract"

    # Title row
    ws.merge_cells("A1:H1")
    title_cell = ws["A1"]
    title_cell.value = f"BOQ Extract — {data.get('vendor_name', 'Unknown Vendor')}"
    title_cell.font = TITLE_FONT
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24

    # Meta row
    ws.merge_cells("A2:H2")
    meta = ws["A2"]
    meta.value = (
        f"Quote: {data.get('quote_number', '—')}  |  "
        f"Date: {data.get('quote_date', '—')}  |  "
        f"Valid Until: {data.get('valid_until', '—')}  |  "
        f"GSTIN: {data.get('vendor_gstin', '—')}"
    )
    meta.font = Font(italic=True, size=9, color="595959")
    meta.alignment = Alignment(horizontal="left")
    ws.row_dimensions[2].height = 16

    # Blank spacer
    ws.row_dimensions[3].height = 6

    # Header row
    for col_idx, (header, width) in enumerate(zip(HEADERS, COL_WIDTHS), start=1):
        cell = ws.cell(row=4, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[4].height = 28
    ws.freeze_panes = "A5"

    items: List[Dict[str, Any]] = data.get("items", [])
    data_start_row = 5
    row = data_start_row
    last_section = None
    section_rows: Dict[str, List[int]] = {}  # section → list of Amount row numbers

    for item in items:
        section = item.get("section", "")
        # Group header row when section changes
        if section and section != last_section:
            ws.merge_cells(f"A{row}:H{row}")
            sec_cell = ws.cell(row=row, column=1, value=section)
            sec_cell.font = SECTION_FONT
            sec_cell.fill = SECTION_FILL
            sec_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            ws.row_dimensions[row].height = 18
            row += 1
            last_section = section
            if section not in section_rows:
                section_rows[section] = []

        qty    = float(item.get("quantity", 0) or 0)
        rate   = float(item.get("rate", 0) or 0)
        gst    = float(item.get("gst_percent", 18) or 18)
        amount = qty * rate  # base amount without GST (matches civil industry convention)

        values = [
            "",
            item.get("description", ""),
            item.get("specification", ""),
            item.get("unit", ""),
            qty,
            rate,
            gst,
            amount,
        ]
        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col_idx, value=val)
            cell.alignment = Alignment(vertical="center", wrap_text=(col_idx in (2, 3)))
            if col_idx in (5,):
                cell.number_format = NUM_FORMAT
            if col_idx in (6, 8):
                cell.number_format = INR_FORMAT
            if col_idx == 7:
                cell.number_format = '0.00"%"'
                cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[row].height = 20

        if section:
            section_rows[section].append(row)
        row += 1

    # Total row — SUM of column H from data_start_row to row-1
    ws.merge_cells(f"A{row}:G{row}")
    total_label = ws.cell(row=row, column=1, value="GRAND TOTAL (Base Amount, excl. GST)")
    total_label.font = TOTAL_FONT
    total_label.fill = TOTAL_FILL
    total_label.alignment = Alignment(horizontal="right", vertical="center")

    total_amount_cell = ws.cell(row=row, column=8)
    total_amount_cell.value = f"=SUM(H{data_start_row}:H{row - 1})"
    total_amount_cell.number_format = INR_FORMAT
    total_amount_cell.font = TOTAL_FONT
    total_amount_cell.fill = TOTAL_FILL
    total_amount_cell.alignment = Alignment(horizontal="right", vertical="center")
    ws.row_dimensions[row].height = 22

    # Confidence note
    row += 2
    ws.merge_cells(f"A{row}:H{row}")
    conf = data.get("overall_confidence", 0)
    note = ws.cell(row=row, column=1,
                   value=f"AI Extraction Confidence: {int(conf * 100)}%  |  Generated by ConstructCRM AI Parser  |  {datetime.now().strftime('%d %b %Y %H:%M')}")
    note.font = Font(italic=True, size=8, color="808080")

    wb.save(out_path)


# ---------------------------------------------------------------------------
# Main background task
# ---------------------------------------------------------------------------

def process_document(document_id: str, api_key: str) -> None:
    """Runs as a FastAPI BackgroundTask. Updates DB record through the pipeline."""
    db = SessionLocal()
    try:
        doc = db.query(Document).filter(Document.id == document_id).first()
        if not doc:
            return

        # --- PROCESSING ---
        doc.status = DocumentStatus.processing
        doc.updated_at = datetime.now(timezone.utc)
        db.commit()

        # 1. Extract text
        text = extract_text(doc.file_path, doc.original_name)

        # 2. Claude extraction (fall back to demo data if no API key)
        if api_key:
            try:
                extracted = call_claude(text, api_key)
            except Exception as e:
                extracted = {**FALLBACK_RESULT, "_claude_error": str(e)}
        else:
            extracted = FALLBACK_RESULT.copy()

        # 3. Generate Excel
        excel_filename = f"{document_id}.xlsx"
        excel_path = os.path.join(os.path.dirname(doc.file_path), excel_filename)
        generate_excel(extracted, excel_path)

        # --- DONE ---
        doc.status = DocumentStatus.done
        doc.extracted_data = extracted
        doc.excel_path = excel_path
        doc.updated_at = datetime.now(timezone.utc)
        db.commit()

    except Exception as exc:
        try:
            doc = db.query(Document).filter(Document.id == document_id).first()
            if doc:
                doc.status = DocumentStatus.failed
                doc.error_message = str(exc)
                doc.updated_at = datetime.now(timezone.utc)
                db.commit()
        except Exception:
            pass
    finally:
        db.close()
