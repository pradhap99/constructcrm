import io
from typing import Dict, Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def _auto_fit(ws, min_w: int = 10, max_w: int = 50) -> None:
    for column in ws.columns:
        max_len = 0
        col_letter = get_column_letter(column[0].column)
        for cell in column:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, min_w), max_w)


def _to_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def export_purchase_order(po_data: Dict[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Purchase Order"

    ws.merge_cells("A1:H1")
    ws["A1"] = "PURCHASE ORDER"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A1"].alignment = Alignment(horizontal="center")

    meta = [
        ("A3", "PO Number:", "B3", po_data.get("po_number", "")),
        ("A4", "Project:", "B4", po_data.get("project", "")),
        ("A5", "Vendor:", "B5", po_data.get("vendor", "")),
        ("A6", "Date:", "B6", str(po_data.get("date", ""))),
        ("E3", "Total Amount (INR):", "F3", po_data.get("total_amount", 0)),
        ("E4", "GST Amount (INR):", "F4", po_data.get("gst_amount", 0)),
    ]
    for label_cell, label, value_cell, value in meta:
        ws[label_cell] = label
        ws[value_cell] = value

    headers = ["#", "Description", "HSN/SAC", "Unit", "Qty", "Rate (INR)", "GST %", "Amount (INR)"]
    row = 8
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center")

    for i, item in enumerate(po_data.get("items", []), 1):
        r = row + i
        for col, val in enumerate([
            i,
            item.get("description", ""),
            item.get("hsn_sac", ""),
            item.get("unit", ""),
            item.get("quantity", 0),
            float(item.get("rate", 0)),
            item.get("gst_percent", 18),
            float(item.get("amount", 0)),
        ], 1):
            ws.cell(row=r, column=col, value=val).border = BORDER

    _auto_fit(ws)
    return _to_bytes(wb)


def export_grn(grn_data: Dict[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "GRN"

    ws.merge_cells("A1:G1")
    ws["A1"] = "GOODS RECEIPT NOTE"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A1"].alignment = Alignment(horizontal="center")

    for label_cell, label, value_cell, value in [
        ("A3", "GRN Number:", "B3", grn_data.get("grn_number", "")),
        ("A4", "PO Number:", "B4", grn_data.get("po_number", "")),
        ("A5", "Vendor:", "B5", grn_data.get("vendor", "")),
        ("A6", "Received Date:", "B6", str(grn_data.get("received_date", ""))),
    ]:
        ws[label_cell] = label
        ws[value_cell] = value

    headers = ["#", "Description", "Unit", "Ordered Qty", "Received Qty", "Accepted Qty", "Remarks"]
    row = 8
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center")

    for i, item in enumerate(grn_data.get("items", []), 1):
        r = row + i
        for col, val in enumerate([
            i,
            item.get("description", ""),
            item.get("unit", ""),
            item.get("ordered_qty", 0),
            item.get("received_qty", 0),
            item.get("accepted_qty", 0),
            item.get("remarks", ""),
        ], 1):
            ws.cell(row=r, column=col, value=val).border = BORDER

    _auto_fit(ws)
    return _to_bytes(wb)


def export_boq(boq_data: Dict[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "BOQ"

    ws.merge_cells("A1:H1")
    ws["A1"] = "BILL OF QUANTITIES"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A1"].alignment = Alignment(horizontal="center")

    for label_cell, label, value_cell, value in [
        ("A3", "BOQ Number:", "B3", boq_data.get("boq_number", "")),
        ("A4", "Project:", "B4", boq_data.get("project", "")),
        ("A5", "Total Amount (INR):", "B5", boq_data.get("total_amount", 0)),
    ]:
        ws[label_cell] = label
        ws[value_cell] = value

    headers = ["#", "Description", "HSN/SAC", "Unit", "Quantity", "Rate (INR)", "Amount (INR)", "Remarks"]
    row = 7
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center")

    for i, item in enumerate(boq_data.get("items", []), 1):
        r = row + i
        for col, val in enumerate([
            item.get("item_no", i),
            item.get("description", ""),
            item.get("hsn_sac_code", ""),
            item.get("unit", ""),
            float(item.get("quantity", 0)),
            float(item.get("rate", 0)),
            float(item.get("amount", 0)),
            item.get("remarks", ""),
        ], 1):
            ws.cell(row=r, column=col, value=val).border = BORDER

    _auto_fit(ws)
    return _to_bytes(wb)


def export_comparative(comp_data: Dict[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Comparative Statement"

    ws.merge_cells("A1:J1")
    ws["A1"] = "COMPARATIVE STATEMENT"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A1"].alignment = Alignment(horizontal="center")

    ws["A3"] = "Comp No:"
    ws["B3"] = comp_data.get("comparative_number", "")
    ws["A4"] = "RFQ Number:"
    ws["B4"] = comp_data.get("rfq_number", "")

    _auto_fit(ws)
    return _to_bytes(wb)
